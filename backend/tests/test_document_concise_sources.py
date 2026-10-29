import asyncio
from copy import deepcopy
from io import BytesIO

import pytest

from app.analysis.markdown import markdown_text
from app.config import Settings
from app.documents.concise import ConciseDocumentGeneration, validate_result
from app.documents.formatting import plain_sources, render_sources
from app.documents.model import DocumentModelClient
from app.documents.tasks import preview, task_config
from .test_note_styles import concise_calls, material


def test_source_links_use_local_labels_and_preserve_headings_and_code():
    markdown = ("## 主题 $x$ [来源](#source-D1-D2)\n\n介绍。\n\n"
                "子主题 [来源](#source-D1)\n---\n\n说明。\n\n"
                "`[来源](#source-D99)`\n\n```md\n## 示例 [来源](#source-D98)\n```\n")
    rendered = render_sources(markdown, {"D1": "幻灯片 [2]", "D2": "幻灯片 6"})
    assert rendered.startswith("## 主题 $x$ [来源：幻灯片 \\[2\\]–幻灯片 6](#document-unit-D1)")
    assert "子主题 [来源：幻灯片 \\[2\\]](#document-unit-D1)\n---" in rendered
    assert "`[来源](#source-D99)`" in rendered
    assert "```md\n## 示例 [来源](#source-D98)\n```" in rendered
    exported = plain_sources(rendered)
    assert "#document-unit-" not in exported
    assert "## 主题 $x$ 来源：幻灯片 \\[2\\]–幻灯片 6" in exported
    assert "幻灯片 [2]" in markdown_text(exported)
    assert "`[来源](#source-D99)`" in exported


@pytest.mark.parametrize("markdown", [
    "只有正文。",
    "## 无来源\n\n说明。",
    "## 主题 [来源](#source-D2)\n\n说明。",
    "## 主题 [来源](#source-D1) [来源](#source-D1)\n\n说明。",
    "## 主题 [来源](#source-D1)\n\n### 缺少来源\n\n说明。",
    "## 主题 [来源](#source-D1)\n\n正文 [来源](#source-D1)。",
    "## 主题 [幻灯片99](#document-unit-D1)\n\n说明。",
    "`[来源](#source-D1)`\n\n```md\n## 主题 [来源](#source-D1)\n```",
    "## 主题 [来源](<#source-D1>)\n\n说明。",
])
def test_missing_unknown_or_misplaced_sources_are_rejected(markdown):
    with pytest.raises(ValueError):
        render_sources(markdown, {"D1": "第 1 页"})


@pytest.mark.parametrize("target", ["D2-D1", "D1-D99"])
def test_ranges_require_real_ordered_endpoints(target):
    with pytest.raises(ValueError):
        render_sources(f"## 主题 [来源](#source-{target})\n\n介绍。", {"D1": "第 8 页", "D2": "第 10 页"})


def test_page_source_shows_real_page_number_instead_of_unit_number():
    assert render_sources("## 主题 [来源](#source-D1)\n\n介绍。", {"D1": "第 8 页"}).startswith(
        "## 主题 [来源：第 8 页](#document-unit-D1)")


def test_presentation_section_shows_actual_slide_range_in_reading_and_exports(client, app, concise_calls):
    from pptx import Presentation
    deck = Presentation()
    for index in range(1, 7):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f"第{index}节"
        slide.placeholders[1].text = "循环的条件与执行过程。"
    stream = BytesIO()
    deck.save(stream)
    concise_calls.hook = lambda payload, value: {**value, "markdown": "## 循环结构 [来源](#source-D2-D6)\n\n本节介绍循环。"}
    client.put("/api/v1/settings/model", json={"token": "fake-slide-range", "stream": False})
    doc = client.post("/api/v1/documents", files={"file": ("来源范围.pptx", stream.getvalue())}).json()
    task = client.post(f'/api/v1/documents/{doc["id"]}/analyses', json={"note_style": "concise"}).json()["task_id"]
    app.state.document_pipeline.run(task)
    result = client.get(f"/api/v1/document-analyses/{task}").json()
    assert "[来源：幻灯片 2–幻灯片 6](#document-unit-D2)" in result["chapters"][0]["markdown"]
    assert len(concise_calls) == 1
    for format in ("md", "txt"):
        text = client.get(f"/api/v1/document-analyses/{task}/notes?format={format}").text
        assert "来源：幻灯片 2–幻灯片 6" in text and "#document-unit-" not in text
    units = [client.get(f'/api/v1/documents/{doc["id"]}/content', params={"offset": index, "limit": 1}).json()["units"][0]
             for index in (1, 5)]
    assert [unit["locator"]["slide"] for unit in units] == [2, 6]


def test_document_source_repair_resume_and_legacy_reading(tmp_path, concise_calls):
    def response(payload, value):
        return {**value, "markdown": "## 主题 [来源](#source-D99)\n\n介绍。" if len(concise_calls) == 1
                else "## 主题 [来源](#source-D1)\n\n介绍。"}
    concise_calls.hook = response

    async def run():
        source = material()
        original = deepcopy(source)
        settings = Settings(llm_key="fake-document-sources", llm_stream=False, note_style="concise")
        async with DocumentModelClient(settings) as model:
            config = task_config(settings, "课程", preview(source, settings, "课程"))
            generation = ConciseDocumentGeneration("sources", tmp_path, config, source, model, asyncio.Event(),
                lambda: None, lambda *args: None, lambda _: None)
            result = await generation.run()
            assert len(concise_calls) == 2
            assert "document_source_must_belong_to_primary" in concise_calls[1][0]["messages"][-1]["content"]
            assert concise_calls[0][1]["primary"] == [{"id": "D1", "text": source["units"][0]["text"]}]
            assert "[来源：正文 1](#document-unit-D1)" in result["chapters"][0]["markdown"]
            assert result["chapters"][0]["source_spans"] == [{"unit_id": "D1", "start_offset": 0,
                                                               "end_offset": len(source["units"][0]["text"])}]
            assert await generation.run() == result and len(concise_calls) == 2
            old = deepcopy(result)
            old["chapters"][0]["markdown"] = "已有成功笔记正文。"
            validate_result(old, source)
            assert source == original
    asyncio.run(run())
