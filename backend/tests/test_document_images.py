"""Document images stay attached to their source; OCR never requires ASR or a remote model."""
import io

import pytest

from app.documents.extraction import extract, EXTRACTION_VERSION
from app.documents.image_text import ImageText
from app.errors import AppError
from app.recognition.media import text_rows


def image_bytes():
    from PIL import Image
    stream = io.BytesIO()
    Image.new("RGB", (120, 60), "white").save(stream, "PNG")
    return stream.getvalue()


def save_document(tmp_path, extension, build):
    path = tmp_path / f"sample.{extension}"
    build(path)
    return path


@pytest.mark.parametrize('kind', ['scatter', 'bubble', 'combination', 'empty'])
@pytest.mark.parametrize('retained', ['', 'title', 'body'])
def test_unreadable_ppt_chart_warns_without_inventing_material(client, app, kind, retained):
    from copy import deepcopy
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData, XyChartData, BubbleChartData
    from pptx.enum.chart import XL_CHART_TYPE
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    if kind in {'scatter', 'bubble'}:
        data = XyChartData() if kind == 'scatter' else BubbleChartData()
        series = data.add_series('测量')
        for x, y in [(1, 10), (2, 20)]:
            series.add_data_point(x, y, *([3] if kind == 'bubble' else []))
        chart_type = XL_CHART_TYPE.XY_SCATTER if kind == 'scatter' else XL_CHART_TYPE.BUBBLE
    else:
        data = CategoryChartData()
        data.categories = ['甲', '乙']
        data.add_series('测量', (None, None) if kind == 'empty' else (1, 2))
        chart_type = XL_CHART_TYPE.COLUMN_CLUSTERED
    chart = slide.shapes.add_chart(chart_type, 0, 0, 4000000, 3000000, data).chart
    if kind == 'combination':
        chart._chartSpace.chart.plotArea.append(deepcopy(chart.plots[0]._element))
    if retained == 'title':
        chart.has_title = True
        chart.chart_title.text_frame.text = '原始标题'
    if retained == 'body':
        slide.shapes.add_textbox(0, 4000000, 3000000, 1000000).text = '独立正文'
    raw = io.BytesIO()
    deck.save(raw)
    response = client.post('/api/v1/documents', files={'file':('chart.pptx', raw.getvalue())})
    assert response.status_code == 201, response.text
    document = response.json()
    result = client.get(f"/api/v1/documents/{document['id']}/content").json()
    text = '\n'.join(u['text'] for u in result['units'])
    assert '图表原始数据' not in text and '图表数据未读取' not in text and '| 分类' not in text
    assert any('幻灯片 1' in warning and '图表' in warning for warning in document['warnings'])
    reason = {'scatter':'散点图', 'bubble':'气泡图', 'combination':'组合图表', 'empty':'分类或数值'}[kind]
    assert any(reason in warning for warning in document['warnings'])
    assert ('原始标题' in text) == (retained == 'title')
    assert ('独立正文' in text) == (retained == 'body')
    preview = client.get(f"/api/v1/documents/{document['id']}/analysis-preview").json()
    assert preview['eligible'] == bool(retained), preview


def test_ppt_category_chart_zero_is_valid_data(tmp_path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    def build(path):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        data = CategoryChartData()
        data.categories = ['甲']
        data.add_series('结果', [0])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, 4000000, 3000000, data)
        deck.save(path)
    content = extract(save_document(tmp_path, 'pptx', build), 'pptx', ImageText())
    assert '| 甲 | 0.0 |' in content['units'][0]['text']
    assert not any('未推断图中结论' in warning for warning in content['warnings'])


def test_ppt_images_group_order_code_superscripts_chart_and_notes(tmp_path, monkeypatch):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    calls = []
    monkeypatch.setattr(ImageText, "recognize", lambda self, raw: (calls.append(raw) or "输出 5050", ""))
    def build(path):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        # Deliberately create objects out of visual order, including a group.
        slide.shapes.add_textbox(0, Inches(4), Inches(3), Inches(1)).text = "下方文字"
        group = slide.shapes.add_group_shape()
        group.shapes.add_picture(io.BytesIO(image_bytes()), 0, Inches(2))
        group.shapes.add_textbox(0, Inches(1), Inches(3), Inches(1)).text = "组内前文"
        slide.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "#1.\ti = 1\n#2.\twhile i < 2:\n#3.\t    i += 1"
        run = slide.shapes.add_textbox(0, Inches(5), Inches(3), Inches(1)).text_frame.paragraphs[0].add_run()
        run.text = "3"
        run._r.get_or_add_rPr().set("baseline", "30000")
        chart = CategoryChartData()
        chart.categories = ["甲", "乙"]
        chart.add_series("人数", (2, 3))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, Inches(6), Inches(3), Inches(2), chart)
        slide.notes_slide.notes_text_frame.text = "讲者的必要补充"
        deck.save(path)
    result = extract(save_document(tmp_path, "pptx", build), "pptx", ImageText())
    assert result["version"] == EXTRACTION_VERSION
    assert len(result["units"]) == 1 and result["units"][0]["locator"] == {"slide": 1}
    text = result["units"][0]["text"]
    assert text.index("组内前文") < text.index("输出 5050") < text.index("下方文字")
    assert "```python\ni = 1\nwhile i < 2:\n    i += 1\n```" in text
    assert "³" in text and "人数" in text and "甲 | 2.0" in text
    assert "讲者备注" in text and "讲者的必要补充" in text and len(calls) == 1


def test_word_images_remain_in_paragraph_and_table_order(tmp_path, monkeypatch):
    from docx import Document
    monkeypatch.setattr(ImageText, "recognize", lambda *_: ("截图中的结果", ""))
    def build(path):
        doc = Document()
        paragraph = doc.add_paragraph("前文")
        paragraph.add_run().add_picture(io.BytesIO(image_bytes()))
        paragraph.add_run("后文")
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0].add_run().add_picture(io.BytesIO(image_bytes()))
        doc.save(path)
    units = extract(save_document(tmp_path, "docx", build), "docx", ImageText())["units"]
    assert units[0]["locator"] == {"block": 1}
    assert units[0]["text"].index("前文") < units[0]["text"].index("截图中的结果") < units[0]["text"].index("后文")
    assert units[1]["locator"] == {"block": 2} and "截图中的结果" in units[1]["text"]


def test_unreadable_slide_image_keeps_other_text(tmp_path):
    from pptx import Presentation
    from pptx.oxml.ns import qn
    def build(path):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = "保留本页正文"
        picture = slide.shapes.add_picture(io.BytesIO(image_bytes()), 0, 1000000)
        picture._pic.blipFill.blip.attrib.pop(qn("r:embed"))
        deck.save(path)
    result = extract(save_document(tmp_path, "pptx", build), "pptx", ImageText())
    assert "保留本页正文" in result["units"][0]["text"]
    assert "图片文字未读取" not in result["units"][0]["text"]
    assert any("幻灯片 1" in warning and "图片数据无法读取" in warning for warning in result["warnings"])
def test_scanned_pdf_reads_image_at_original_page(tmp_path, monkeypatch):
    from .pdf_samples import PdfSample
    monkeypatch.setattr(ImageText, "recognize", lambda *_: ("扫描页的文字", ""))
    def build(path):
        document = PdfSample()
        document.page(100, 100)
        document.image(document.page(100, 100), (0, 0, 100, 100), image_bytes())
        document.save(path)
    units = extract(save_document(tmp_path, "pdf", build), "pdf", ImageText())["units"]
    assert units[0]["text"] == "" and units[1]["locator"] == {"page": 2}
    assert "扫描页的文字" in units[1]["text"] and "OCR" in units[1]["text"]


def test_pdf_ocr_failure_saves_report_without_losing_page_or_leaking_error(client, app, tmp_path, monkeypatch):
    import subprocess
    from PIL import Image
    def fail(*args, **kwargs):
        raise subprocess.CalledProcessError(1, ["isolated-ocr"], stderr=b"isolated OCR backend detail")
    monkeypatch.setattr("app.recognition.media.Recognizer.ocr", fail)
    stream = io.BytesIO()
    Image.new("RGB", (120, 80), "white").save(stream, "PDF")
    response = client.post("/api/v1/documents", files={"file": ("scanned.pdf", stream.getvalue())})
    assert response.status_code == 201, response.text
    warnings = response.json()["warnings"]
    assert any("第 1 页 · 图片1" in warning and "错误编号" in warning and "log/" in warning for warning in warnings)
    assert "isolated OCR backend detail" not in str(warnings)
    reports = list(app.state.pipeline.get_settings().log_dir.glob("*.log"))
    assert len(reports) == 1
    assert "document_ocr_failed" in reports[0].read_text(encoding="utf-8")
    assert "isolated OCR backend detail" in reports[0].read_text(encoding="utf-8")
    content = client.get(f'/api/v1/documents/{response.json()["id"]}/content').json()
    assert content["units"][0]["locator"] == {"page": 1}
    assert content["units"][0]["text"] == ""
    assert client.get(f'/api/v1/documents/{response.json()["id"]}/analysis-preview').json()["eligible"] is False


def test_duplicate_images_only_run_ocr_once_but_keep_each_source(monkeypatch):
    calls = []
    monkeypatch.setattr(ImageText, "recognize", lambda self, raw: (calls.append(raw) or "重复使用的图片", ""))
    images = ImageText()
    assert "重复使用的图片" in images.read(b"same", "幻灯片1")
    assert "重复使用的图片" in images.read(b"same", "幻灯片2")
    assert len(calls) == 1


def test_ocr_fragments_keep_rows_and_left_to_right_order():
    boxes = [[[20, 0], [30, 0], [30, 10], [20, 10]],
             [[0, 1], [10, 1], [10, 11], [0, 11]],
             [[0, 20], [50, 20], [50, 30], [0, 30]]]
    assert text_rows(boxes, ["4", "1", "[1, 2, 3]"]) == "1 4\n[1, 2, 3]"


def test_ocr_limits_and_failures_do_not_leak_engine_errors(tmp_path, monkeypatch):
    from app.config import Settings
    images = ImageText(Settings(model_dir=tmp_path / "models"))
    monkeypatch.setattr("app.documents.image_text.MAX_IMAGE_PIXELS", 1)
    with pytest.raises(AppError, match="像素"):
        images.read(image_bytes(), "第一张")
    assert images.read(b"not an image", "第二张") == ""
    assert "第二张" in images.warnings[0] and "读取失败" in images.warnings[0]
    assert "not an image" not in str(images.warnings)


def test_numbered_code_inside_text_and_non_code_numbering(tmp_path):
    from pptx import Presentation
    def build(path):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = "例题说明\n#1.\tx = 1\n#2.\tprint(x)"
        slide.shapes.add_textbox(0, 2000000, 1000000, 1000000).text = "#1. 课程目录\n#2. 课程介绍"
        deck.save(path)
    text = extract(save_document(tmp_path, "pptx", build), "pptx")["units"][0]["text"]
    assert "例题说明\n```python\nx = 1\nprint(x)\n```" in text
    assert text.count("```python") == 1 and "课程目录" in text


@pytest.mark.parametrize("problem", ["识别失败", "没有识别到文字"])
def test_failed_image_is_not_silently_treated_as_absent(monkeypatch, problem):
    monkeypatch.setattr(ImageText, "recognize", lambda *_: ("", problem))
    images = ImageText()
    assert images.read(b"image", "幻灯片17") == ""
    assert images.warnings == [f"幻灯片17：{problem}"]


def test_no_ocr_text_never_becomes_source_material(client, monkeypatch):
    from PIL import Image
    calls = []
    monkeypatch.setattr(ImageText, "recognize", lambda self, raw: (calls.append(raw) or "", "未识别到文字，可能为非文字图片或识别遗漏。"))
    stream = io.BytesIO()
    Image.new("RGB", (120, 80), "white").save(stream, "PDF")
    response = client.post("/api/v1/documents", files={"file": ("empty-image.pdf", stream.getvalue())})
    assert response.status_code == 201
    value = response.json()
    assert value["character_count"] == 0 and len(calls) == 1
    assert any("未识别到文字" in warning for warning in value["warnings"])
    content = client.get(f'/api/v1/documents/{value["id"]}/content').json()
    assert content["units"][0]["text"] == ""
    assert content["units"][0]["locator"] == {"page": 1}
    assert client.get(f'/api/v1/documents/{value["id"]}/analysis-preview').json()["eligible"] is False


def test_image_count_limit_is_not_silent_truncation(monkeypatch):
    monkeypatch.setattr("app.documents.image_text.MAX_IMAGES", 1)
    monkeypatch.setattr(ImageText, "recognize", lambda *_: ("文本", ""))
    images = ImageText()
    images.read(b"one", "第一张")
    with pytest.raises(AppError, match="拆分"):
        images.read(b"two", "第二张")


def test_upload_uses_local_ocr_without_initializing_asr(client, monkeypatch):
    from pptx import Presentation
    calls = []
    def ocr(self, path, **kwargs):
        assert self.model is None and self.chinese is None
        assert kwargs == {"preserve_rows": True}
        calls.append(path.read_bytes())
        return "本地识别结果", 95
    monkeypatch.setattr("app.recognition.media.Recognizer.ocr", ocr)
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_picture(io.BytesIO(image_bytes()), 0, 0)
    stream = io.BytesIO()
    deck.save(stream)
    response = client.post("/api/v1/documents", files={"file": ("截图.pptx", stream.getvalue())})
    assert response.status_code == 201
    content = client.get(f'/api/v1/documents/{response.json()["id"]}/content').json()
    assert "本地识别结果" in content["units"][0]["text"] and len(calls) == 1


def test_ocr_does_not_hold_task_control_lock(client, app, monkeypatch):
    observed = []
    from app.documents import store as service
    original = service.extract
    def extraction(*args):
        acquired = app.state.worker.lock.acquire(blocking=False)
        observed.append(acquired)
        if acquired:
            app.state.worker.lock.release()
        return original(*args)
    monkeypatch.setattr(service, "extract", extraction)
    response = client.post("/api/v1/documents", files={"file": ("文字.txt", b"example")})
    assert response.status_code == 201 and observed == [True]


def test_concurrent_duplicate_uploads_recheck_after_extraction(client, app, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    from app.documents import store as service
    original, barrier = service.extract, Barrier(2)
    def extraction(*args):
        barrier.wait(timeout=5)
        return original(*args)
    monkeypatch.setattr(service, "extract", extraction)
    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(lambda _: client.post("/api/v1/documents", files={"file": ("相同.txt", b"same material")}), range(2)))
    assert all(response.status_code == 201 for response in responses)
    assert len({response.json()["id"] for response in responses}) == 1
    assert sorted(response.json()["reused"] for response in responses) == [False, True]
    assert len(list((app.state.pipeline.get_settings().data_dir / "documents").iterdir())) == 1


def test_existing_extraction_is_not_overwritten_by_new_parser(client, app):
    import json
    response = client.post("/api/v1/documents", files={"file": ("历史.txt", b"material")})
    identifier = response.json()["id"]
    path = app.state.pipeline.get_settings().data_dir / "documents" / identifier / "content.json"
    value = json.loads(path.read_text("utf-8"))
    value["version"] = 1
    path.write_text(json.dumps(value), "utf-8")
    before = path.read_bytes()
    repeated = client.post("/api/v1/documents", files={"file": ("历史.txt", b"material")})
    assert repeated.json()["reused"] and repeated.json()["id"] == identifier
    assert path.read_bytes() == before
