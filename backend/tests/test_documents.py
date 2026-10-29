import io
import zipfile

import pytest
from sqlalchemy import select

from app.db import Document, Task, Video
from app.documents.extraction import DOCUMENT_EXTENSIONS, extract
from app.errors import AppError
from .document_samples import sample


@pytest.mark.parametrize("extension", DOCUMENT_EXTENSIONS)
def test_document_intake_lifecycle(client, app, extension):
    raw = sample(extension)
    response = client.post("/api/v1/documents", files={"file": (f"lesson.{extension}", raw)})
    assert response.status_code == 201, response.text
    document = response.json()
    assert document["kind"] == "document"
    assert document["unit_count"] > 0 and document["character_count"] > 0
    assert "fingerprint" not in document
    identifier = document["id"]
    assert client.get("/api/v1/documents").json() == [{key: value for key, value in document.items() if key != "reused"}]
    units = client.get(f"/api/v1/documents/{identifier}/content?limit=1").json()
    assert len(units["units"]) == 1
    assert units["units"][0]["locator"]
    assert client.get(f"/api/v1/documents/{identifier}/original").content == raw
    again = client.post("/api/v1/documents", files={"file": (f"other.{extension}", raw)}).json()
    assert again["reused"] and again["id"] == identifier
    assert len(list((app.state.pipeline.get_settings().data_dir / "documents").iterdir())) == 1
    assert client.patch(f"/api/v1/documents/{identifier}", json={"filename": "renamed", "pinned": True}).json()["pinned"]
    # Editing a display name never changes the original's format or content.
    assert client.get(f"/api/v1/documents/{identifier}/original").content == raw
    with app.state.sessions() as session:
        assert session.scalar(select(Task)) is None
        assert session.scalar(select(Video)) is None
    assert client.delete(f"/api/v1/documents/{identifier}").json()["deleted"]
    assert client.get(f"/api/v1/documents/{identifier}/content").status_code == 404
    assert not list((app.state.pipeline.get_settings().data_dir / "documents").iterdir())


def extracted(tmp_path, extension, data=None):
    path = tmp_path / f"sample.{extension}"
    path.write_bytes(data if data is not None else sample(extension))
    return extract(path, extension)


def test_source_locations_and_order(tmp_path):
    pdf = extracted(tmp_path, "pdf")["units"]
    assert [unit["locator"] for unit in pdf] == [{"page": 1}, {"page": 2}]
    assert "Second page evidence" in pdf[1]["text"]
    word = extracted(tmp_path, "docx")["units"]
    assert [unit["locator"]["block"] for unit in word] == [1, 2, 3, 4]
    assert word[1]["text"] == "表格之前" and word[3]["text"] == "表格之后"
    assert "样本" in word[2]["text"] and word[2]["format"] == "markdown"
    ppt = extracted(tmp_path, "pptx")["units"]
    assert ppt[1]["locator"] == {"slide": 2}
    assert "第二张幻灯片" in ppt[1]["text"]
    excel = extracted(tmp_path, "xlsx")["units"]
    assert excel[0]["locator"] == {"sheet": "测量记录", "row_start": 1, "row_end": 3}
    assert r"=B2\*2" in excel[0]["text"] and "第二张表" in excel[1]["label"]
    csv = extracted(tmp_path, "csv")["units"]
    assert csv[0]["locator"] == {"record_start": 1, "record_end": 2}
    assert "第一行 / 第二行" in csv[0]["text"]


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "gb18030"])
def test_text_encodings_and_source_lines(tmp_path, encoding):
    value = extracted(tmp_path, "txt", "中文\r\n第二行".encode(encoding))
    assert value["units"][0]["text"] == "中文\n第二行"
    assert value["units"][0]["locator"] == {"line_start": 1, "line_end": 2}


def test_large_text_pagination(client):
    raw = "\n".join(f"line{i}" for i in range(3100)).encode()
    doc = client.post("/api/v1/documents", files={"file": ("long.txt", raw)}).json()
    page = client.get(f'/api/v1/documents/{doc["id"]}/content?offset=30').json()
    assert page["total"] == 31
    assert page["units"][0]["locator"] == {"line_start": 3001, "line_end": 3100}
    assert client.get(f'/api/v1/documents/{doc["id"]}/content?offset=-1').status_code == 422


def test_blank_scanned_and_encrypted_pdf(tmp_path):
    from .pdf_samples import PdfSample
    document = PdfSample()
    document.page(100, 100)
    parsed = extracted(tmp_path, "pdf", document.bytes())
    assert parsed["units"][0]["text"] == ""
    assert any("可能是空白页或当前无法读取的内容" in warning for warning in parsed["warnings"])
    document.writer.encrypt("private", "owner", algorithm="AES-256")
    encrypted = document.bytes()
    with pytest.raises(AppError, match="加密"):
        extracted(tmp_path, "pdf", encrypted)


@pytest.mark.parametrize("filename,raw", [("fake.pdf", b"bad"), ("fake.docx", b"bad"), ("old.doc", b"bad"), ("binary.txt", b"\x00"), ("empty.txt", b"")])
def test_invalid_upload_leaves_no_record_or_folder(client, app, filename, raw):
    response = client.post("/api/v1/documents", files={"file": (filename, raw)})
    assert response.status_code == 400
    assert client.get("/api/v1/documents").json() == []
    assert not list((app.state.pipeline.get_settings().data_dir / "documents").iterdir())


def test_extract_failure_cleanup(client, app, monkeypatch):
    def fail(*args):
        raise RuntimeError("simulated failure")
    monkeypatch.setattr("app.documents.store.extract", fail)
    assert client.post("/api/v1/documents", files={"file": ("a.txt", b"text")}).status_code == 500
    assert client.get("/api/v1/documents").json() == []
    assert not list((app.state.pipeline.get_settings().data_dir / "documents").iterdir())


def test_database_failure_preserves_disk_consistency(client, app):
    from sqlalchemy import event
    def fail(*args):
        raise RuntimeError("simulated database failure")
    event.listen(Document, "before_insert", fail)
    try:
        assert client.post("/api/v1/documents", files={"file": ("a.txt", b"text")}).status_code == 500
    finally:
        event.remove(Document, "before_insert", fail)
    assert client.get("/api/v1/documents").json() == []
    assert not list((app.state.pipeline.get_settings().data_dir / "documents").iterdir())
    document = client.post("/api/v1/documents", files={"file": ("a.txt", b"text")}).json()
    event.listen(Document, "before_delete", fail)
    try:
        assert client.delete(f'/api/v1/documents/{document["id"]}').status_code == 500
    finally:
        event.remove(Document, "before_delete", fail)
    assert client.get(f'/api/v1/documents/{document["id"]}/original').content == b"text"


def test_literal_table_cells_and_markdown_source(tmp_path):
    unit = extracted(tmp_path, "csv", b'column\n"<img src=x> $a$ [text](https://example.invalid)"')["units"][0]
    assert r"\<img" in unit["text"] and r"\$a\$" in unit["text"]
    assert r"\[text\]" in unit["text"]
    assert "<script>" in extracted(tmp_path, "md")["units"][0]["text"]  # renderer enforces safety, extraction preserves source


def test_archive_size_bound_and_no_archive_extraction(tmp_path, monkeypatch):
    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", "xml")
        archive.writestr("../../outside.txt", "untrusted")
    def oversized(self):
        item = zipfile.ZipInfo("word/document.xml")
        item.file_size = 129 * 1024 * 1024
        return [item]
    monkeypatch.setattr(zipfile.ZipFile, "infolist", oversized)
    with pytest.raises(AppError, match="解压后过大"):
        extract(path, "docx")
    assert not (tmp_path.parent / "outside.txt").exists()


def test_document_request_limit_and_paths(client, monkeypatch):
    monkeypatch.setattr("app.http.MAX_DOCUMENT_BYTES", 10)
    response = client.post("/api/v1/documents", content=b"", headers={"content-length": str(2 * 1024 * 1024)})
    assert response.status_code == 413
    assert client.get("/api/v1/documents/missing/content").status_code == 404
    assert client.get("/api/v1/health").json()["supported_document_extensions"] == list(DOCUMENT_EXTENSIONS)


def test_reopen_database_reads_saved_document(client, app):
    from app.config import Settings
    from app.main import create_app
    from fastapi.testclient import TestClient
    doc = client.post("/api/v1/documents", files={"file": ("saved.md", sample("md"))}).json()
    settings = app.state.pipeline.get_settings()
    reopened = create_app(Settings(data_dir=settings.data_dir, model_dir=settings.model_dir, log_dir=settings.log_dir), start_worker=False)
    with TestClient(reopened) as other:
        content = other.get(f'/api/v1/documents/{doc["id"]}/content').json()
        assert "**重点**" in content["units"][0]["text"]


@pytest.mark.parametrize("extension", ["pdf", "docx", "pptx"])
def test_import_ocr_is_per_material_and_reuses_identical_input(client, app, monkeypatch, extension):
    import json
    from app.documents.material import snapshot
    monkeypatch.setattr("app.documents.store.runtime_ready", lambda _: True)
    monkeypatch.setattr("app.documents.store.model_ready", lambda *_: True)
    calls = []
    def capture(path, kind, reader):
        calls.append(reader.settings.ocr_engine)
        return extract(path, kind, reader)
    monkeypatch.setattr("app.documents.store.extract", capture)
    current = app.state.settings_service.current
    # Office archives contain timestamps: use identical bytes for every import.
    raw = sample(extension)
    def upload(engine):
        response = client.post("/api/v1/documents", data={"ocr_option": engine},
                               files={"file": (f"lesson.{extension}", raw)})
        assert response.status_code == 201, response.text
        return response.json()
    first = upload("rapidocr")
    path = app.state.document_analyses.documents.folder(first["id"]) / "content.json"
    saved = path.read_bytes()
    assert upload("rapidocr")["id"] == first["id"]
    assert calls == ["rapidocr"]
    assert path.read_bytes() == saved and json.loads(saved)["ocr_option"] == "rapidocr"
    content = json.loads(saved)
    without_ocr = {k: v for k, v in content.items() if k != "ocr_option"}
    assert snapshot(content) == snapshot(without_ocr)
    assert app.state.settings_service.current is current
    assert len(client.get("/api/v1/documents").json()) == 1


@pytest.mark.parametrize("runtime,files,code", [(False, True, "runtime_missing"), (True, False, "model_not_downloaded")])
def test_import_unavailable_ocr_does_not_extract_or_register(client, app, monkeypatch, runtime, files, code):
    monkeypatch.setattr("app.documents.store.runtime_ready", lambda _: runtime)
    monkeypatch.setattr("app.documents.store.model_ready", lambda *_: files)
    monkeypatch.setattr("app.documents.store.extract", lambda *_: pytest.fail("must not extract"))
    response = client.post("/api/v1/documents", data={"ocr_option": "rapidocr"}, files={"file": ("lesson.pdf", sample("pdf"))})
    assert response.status_code == 400 and response.json()["error"]["code"] == code
    assert client.get("/api/v1/documents").json() == []
    assert not list((app.state.pipeline.get_settings().data_dir / "documents").iterdir())


def test_import_plain_text_does_not_need_ocr_and_rejects_unknown_selection(client, monkeypatch):
    monkeypatch.setattr("app.documents.store.runtime_ready", lambda _: pytest.fail("text has no OCR"))
    assert client.post("/api/v1/documents", data={"ocr_option": "rapidocr"}, files={"file": ("a.txt", b"text")}).status_code == 201
    assert client.post("/api/v1/documents", data={"ocr_option": "arbitrary"}, files={"file": ("a.pdf", b"invalid")}).status_code == 422


def test_import_freezes_ocr_during_extraction(client, app, monkeypatch):
    from dataclasses import replace
    from pptx import Presentation
    from .test_document_images import image_bytes
    from app.documents.image_text import ImageText
    settings = app.state.settings_service
    monkeypatch.setattr("app.documents.store.runtime_ready", lambda _: True)
    monkeypatch.setattr("app.documents.store.model_ready", lambda *_: True)
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_picture(io.BytesIO(image_bytes()), 0, 0)
    stream = io.BytesIO()
    deck.save(stream)
    calls = []
    def recognize(reader, raw):
        app.state.current_settings.replace(replace(settings.current, ocr_min_confidence=0.9))
        calls.append(reader.settings.ocr_min_confidence)
        return "所选模型读取的图片文字", ""
    monkeypatch.setattr(ImageText, "recognize", recognize)
    result = client.post("/api/v1/documents", data={"ocr_option": "rapidocr"}, files={"file": ("image.pptx", stream.getvalue())})
    assert result.status_code == 201, result.text
    assert calls == [0.65]
    assert "所选模型读取的图片文字" in client.get(f'/api/v1/documents/{result.json()["id"]}/content').json()["units"][0]["text"]
