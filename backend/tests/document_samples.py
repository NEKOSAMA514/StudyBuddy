"""Small generated documents, shared by parser and browser tests."""
from io import BytesIO


def sample(extension):
    stream = BytesIO()
    if extension == "pdf":
        from .pdf_samples import text_pdf
        stream.write(text_pdf(("First page evidence", "Second page evidence")))
    elif extension == "docx":
        from docx import Document
        doc = Document()
        doc.add_heading("学习材料", 1)
        doc.add_paragraph("表格之前")
        table = doc.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, [("名称", "数值"), ("样本", "12")]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        doc.add_paragraph("表格之后")
        doc.save(stream)
    elif extension == "pptx":
        from pptx import Presentation
        slides = Presentation()
        for title in ("第一张幻灯片", "第二张幻灯片"):
            slide = slides.slides.add_slide(slides.slide_layouts[1])
            slide.shapes.title.text = title
            slide.placeholders[1].text = "需要保留的正文"
        slides.save(stream)
    elif extension == "xlsx":
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.title = "测量记录"
        workbook.active.append(["名称", "数值"])
        workbook.active.append(["样本", 12])
        workbook.active.append(["计算", "=B2*2"])
        other = workbook.create_sheet("第二张表")
        other.append(["不同内容"])
        workbook.save(stream)
        workbook.close()
    else:
        return {"csv": '名称,说明\n样本,"第一行\n第二行"\n', "txt": "第一行\n第二行", "md": "# 学习材料\n\n**重点**\n\n<script>alert(1)</script>\n\n![remote](https://example.invalid/image.png)"}[extension].encode("utf-8")
    return stream.getvalue()
