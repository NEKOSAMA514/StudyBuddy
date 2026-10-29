"""提取可读正文与稳定来源位置，不调用生成模型。"""
import csv
import io
import re
import zipfile
from pathlib import Path

from ..errors import AppError, TaskInterrupted
from .formatting import literal, table

DOCUMENT_EXTENSIONS = ("pdf", "docx", "pptx", "xlsx", "csv", "txt", "md")
MAX_DOCUMENT_BYTES = 100 * 1024 * 1024
MAX_TEXT = 5_000_000
MAX_UNITS = 10_000
EXTRACTION_VERSION = 2


def decode_text(path):
    if path.stat().st_size > 16 * 1024 * 1024:
        raise AppError("document_limit", "文本文件过大，当前最多解析 16 MB 文本。")
    raw = path.read_bytes()
    encodings = ["utf-16"] if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else ["utf-8-sig", "gb18030"]
    for encoding in encodings:
        try:
            value = raw.decode(encoding)
            if "\x00" not in value:
                return value.replace("\r\n", "\n").replace("\r", "\n"), encoding
        except UnicodeError:
            pass
    raise AppError("document_encoding", "无法读取文本编码，请另存为 UTF-8 后上传。")


def validate_archive(path, extension):
    required = {"docx": "word/document.xml", "pptx": "ppt/presentation.xml", "xlsx": "xl/workbook.xml"}
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 20_000 or sum(item.file_size for item in entries) > 128 * 1024 * 1024:
            raise AppError("document_limit", "文档解压后过大，当前最多解析 128 MB 内容。")
        if required[extension] not in archive.namelist():
            raise AppError("invalid_document", "文件内容与扩展名不符，请上传正确的文档。")


# 按文件格式分派本地提取，并统一约束来源块数量、文字量及输出结构。
def extract(path: Path, extension: str, images=None):
    from .image_text import ImageText
    images = images if images is not None else ImageText()
    units, warnings = [], []
    count = 0
    check = getattr(images, "check", lambda: None)

    def add(text, label, locator, markdown=False):
        nonlocal count
        check()
        count += len(text)
        if count > MAX_TEXT or len(units) >= MAX_UNITS:
            raise AppError("document_limit", "文档内容过多，当前最多解析 500 万字符、10,000 个内容块，请拆分后上传。")
        units.append({"id": f"D{len(units) + 1}", "label": label, "locator": locator,
                      "text": text, "format": "markdown" if markdown else "text"})

    try:
        if extension in {"docx", "pptx", "xlsx"}:
            validate_archive(path, extension)
        if extension == "pdf":
            from .pdf import PdfDocument, page_images
            with PdfDocument(path, check=check, max_pages=MAX_UNITS) as reader:
                for index in range(reader.page_count):
                    page = reader.read_page(index)
                    _, records = page_images(page, images)
                    parts = [page["text"]]
                    images.warnings.extend(page["warnings"])
                    parts.extend(record["text"] for record in records)
                    add("\n\n".join(p for p in parts if p.strip()), f"第 {index + 1} 页", {"page": index + 1})
            empty = sum(not unit["text"].strip() for unit in units)
            if empty:
                warnings.append(f"{empty} 页没有可提取的文字，可能是空白页或当前无法读取的内容。")
            warnings.append("PDF 提取文本层并对页面实际绘制的图片进行本地 OCR；图片文字单独标注，页内图片附在文本之后。分栏、表格、矢量文字及公式结构可能不完整。")
        elif extension == "docx":
            from .office import word_units
            for text, label, locator, markdown in word_units(path, images):
                add(text, label, locator, markdown)
            warnings.append("Word 按正文顺序提取段落、表格及内嵌图片文字，不推算页码；图片文字来自本地 OCR。页眉页脚、批注、修订及复杂公式未提取。")
        elif extension == "pptx":
            from .office import presentation_units
            for text, label, locator in presentation_units(path, images):
                add(text, label, locator, True)
            warnings.append("PPT 按位置提取文本框、组合形状、表格、普通分类图表数据、讲者备注和图片文字；图片文字来自本地 OCR。动画、SmartArt、复杂公式及部分图表类型无法完整读取。")
        elif extension == "xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                for sheet in workbook:
                    sheet.reset_dimensions()
                    batch, first = [], 1
                    for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        if number > 100_000 or len(row) > 1000:
                            raise AppError("document_limit", "工作表超过 100,000 行或 1,000 列，请拆分后上传。")
                        batch.append(row)
                        if len(batch) == 100:
                            add(table(batch), f"{sheet.title} · 第 {first}–{number} 行",
                                {"sheet": sheet.title, "row_start": first, "row_end": number}, True)
                            batch, first = [], number + 1
                    if batch:
                        add(table(batch), f"{sheet.title} · 第 {first}–{number} 行",
                            {"sheet": sheet.title, "row_start": first, "row_end": number}, True)
            finally:
                workbook.close()
            warnings.append("Excel 保留单元格内容和公式原文，不计算公式；不提取图表、图片或样式。表头数字表示列序号。")
        else:
            value, encoding = decode_text(path)
            if extension == "csv":
                try:
                    dialect = csv.Sniffer().sniff(value[:8192], delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                batch, first, number = [], 1, 0
                for number, row in enumerate(csv.reader(io.StringIO(value), dialect), 1):
                    if len(row) > 1000:
                        raise AppError("document_limit", "CSV 超过 1,000 列，请拆分后上传。")
                    batch.append(row)
                    if len(batch) == 100:
                        add(table(batch), f"第 {first}–{number} 条记录", {"record_start": first, "record_end": number}, True)
                        batch, first = [], number + 1
                if batch:
                    add(table(batch), f"第 {first}–{number} 条记录", {"record_start": first, "record_end": number}, True)
                warnings.append("CSV 按逻辑记录编号（含换行的单元格仍算一条）；表头数字表示列序号。")
            else:
                lines = value.splitlines()
                for start in range(0, len(lines), 100):
                    end = min(start + 100, len(lines))
                    add("\n".join(lines[start:end]), f"第 {start + 1}–{end} 行", {"line_start": start + 1, "line_end": end}, extension == "md")
                if extension == "md":
                    warnings.append("Markdown 按 100 行分块预览，跨块的代码或表格可能显示不完整；不加载图片或执行 HTML。")
            if encoding != "utf-8-sig":
                warnings.append(f"已按 {encoding} 编码读取文本。")
        if not any(unit["text"].strip() for unit in units):
            warnings.append("未提取到可读文字；原文件已保留。")
        return {"version": EXTRACTION_VERSION, "units": units, "warnings": warnings + images.warnings, "character_count": count}
    except (AppError, TaskInterrupted):
        raise
    except ImportError as exc:
        raise AppError("document_dependency", "文档解析依赖未安装，请安装 backend/requirements.txt 后重试。", 503) from exc
    except Exception as exc:
        raise AppError("invalid_document", "文档无法解析，可能已损坏、加密或与扩展名不符。") from exc
