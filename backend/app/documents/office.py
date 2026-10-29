"""按原顺序读取 Office 内容，将图片文字保留在对应来源旁。"""
import re

from .formatting import literal, table
from ..errors import AppError


# 只将连续编号且具有代码特征的文本恢复为代码块，避免误判普通列表。
def numbered_code(text):
    lines = text.splitlines()
    matches = [re.fullmatch(r"#(\d+)\.[\t ]?(.*)", line) for line in lines]
    if len(lines) < 2 or not all(matches) or [int(m[1]) for m in matches] != list(range(1, len(lines) + 1)):
        return None
    code = "\n".join(m[2] for m in matches)
    if not re.search(r"(?m)^\s*(?:for |while |if |import |def |print\(|[A-Za-z_]\w*\s*=)", code):
        return None
    fence = "`" * max(3, max((len(m[0]) + 1 for m in re.finditer(r"`+", code)), default=0))
    return f"{fence}python\n{code}\n{fence}"


# 按段落和文本片段读取幻灯片，保留换行及数字上下标。
def slide_text(frame):
    paragraphs = []
    superscript = str.maketrans("0123456789+-=()", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾")
    subscript = str.maketrans("0123456789+-=()", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎")
    for paragraph in frame.paragraphs:
        parts = []
        for child in paragraph._p:
            if child.tag.endswith("}br"):
                parts.append("\n")
            elif child.tag.endswith("}r") or child.tag.endswith("}fld"):
                text = "".join(node.text or "" for node in child.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t"))
                properties = child.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
                baseline = int(properties.get("baseline", "0")) if properties is not None else 0
                parts.append(text.translate(superscript if baseline > 0 else subscript) if baseline else text)
        paragraphs.append("".join(parts))
    text = "\n".join(paragraphs)
    lines, parts, index = text.splitlines(), [], 0
    while index < len(lines):
        end = index
        while end < len(lines) and re.match(r"^#\d+\.", lines[end]):
            end += 1
        if end > index:
            group = "\n".join(lines[index:end])
            parts.append(numbered_code(group) or literal(group))
            index = end
        else:
            parts.append(literal(lines[index]))
            index += 1
    return "\n".join(parts)


# 按幻灯片位置读取正文、图片、备注及支持的图表，未读数据仅记警告。
def presentation_units(path, images):
    from pptx import Presentation
    from pptx.enum.chart import XL_CHART_TYPE

    def shape_text(shape, label):
        if shape.shape_type == 6:
            return "\n\n".join(shape_text(child, label) for child in sorted(shape.shapes, key=lambda item: (item.top, item.left)))
        if shape.has_table:
            return table([[cell.text for cell in row.cells] for row in shape.table.rows])
        if shape.has_chart:
            chart = shape.chart
            title = ""
            reason = "图表数据无法完整读取"
            try:
                title = literal(chart.chart_title.text_frame.text) if chart.has_title and chart.chart_title.has_text_frame else ""
                if len(chart.plots) != 1:
                    reason = "组合图表暂不支持读取数据"
                    raise ValueError("combined_chart")
                if chart.chart_type in {XL_CHART_TYPE.XY_SCATTER, XL_CHART_TYPE.XY_SCATTER_LINES,
                        XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, XL_CHART_TYPE.XY_SCATTER_SMOOTH,
                        XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS, XL_CHART_TYPE.BUBBLE, XL_CHART_TYPE.BUBBLE_THREE_D_EFFECT}:
                    reason = "图表为散点图或气泡图，暂不支持读取数据"
                    raise ValueError("non_category_chart")
                categories = [category.label for category in chart.plots[0].categories]
                series = list(chart.series)
                values = [list(item.values) for item in series]
                if (not categories or not series or any(len(row) != len(categories) for row in values)
                        or not any(value is not None for row in values for value in row)):
                    reason = "图表缺少可完整读取的分类或数值"
                    raise ValueError("missing_category_data")
                rows = [["分类", *[item.name for item in series]]]
                rows.extend([name, *[row[i] for row in values]] for i, name in enumerate(categories))
                return title + "\n\n图表原始数据：\n\n" + table(rows)
            except (AttributeError, ValueError, IndexError):
                images.warnings.append(f"{label}：{reason}，未推断图中结论。")
                return title
        if shape.shape_type == 13:
            image_label = f"{label} · 图片{shape.shape_id}"
            try:
                return literal(images.read(shape.image.blob, image_label))
            except AppError:
                raise
            except Exception:
                images.warnings.append(f"{image_label}：图片数据无法读取，未识别其中内容。")
                return ""
        return slide_text(shape.text_frame) if shape.has_text_frame else ""

    for number, slide in enumerate(Presentation(path).slides, 1):
        label = f"幻灯片 {number}"
        parts = [shape_text(shape, label) for shape in sorted(slide.shapes, key=lambda item: (item.top, item.left))]
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None and frame.text.strip():
                parts.append("讲者备注：\n" + literal(frame.text))
        yield "\n\n".join(part for part in parts if part.strip()), label, {"slide": number}


# 沿 Word 原始块顺序提取段落和表格，内嵌图片文字保留在所属位置。
def word_units(path, images):
    from docx import Document
    from docx.table import Table
    from docx.oxml.ns import qn

    def paragraph_text(paragraph, label):
        parts = []
        for node in paragraph._p.iter():
            if node.tag == qn("w:t"):
                parts.append(node.text or "")
            elif node.tag in {qn("w:br"), qn("w:cr")}:
                parts.append("\n")
            elif node.tag == qn("w:tab"):
                parts.append("\t")
            elif node.tag == qn("a:blip"):
                identifier = node.get(qn("r:embed"))
                if identifier:
                    part = paragraph.part.related_parts[identifier]
                    parts.append("\n" + images.read(part.blob, label) + "\n")
                else:
                    images.warnings.append(f"{label}：外部链接图片未读取。")
                    parts.append("\n外部链接图片未读取。\n")
        return "".join(parts)

    def block_text(block, label):
        if isinstance(block, Table):
            return table([["\n".join(block_text(item, label) for item in cell.iter_inner_content()) for cell in row.cells]
                          for row in block.rows])
        return paragraph_text(block, label)

    for number, block in enumerate(Document(path).iter_inner_content(), 1):
        is_table = isinstance(block, Table)
        label = f"正文块 {number}" + (" · 表格" if is_table else "")
        text = block_text(block, label)
        if text.strip():
            yield text, label, {"block": number}, is_table
